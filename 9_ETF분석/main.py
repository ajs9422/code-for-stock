# =============================================================================
# ETF 가치 분석 스캐너
#
# [개요]
# 단순 수익률이 아닌, ETF 구성 종목의 S-RIM 적정주가 대비 저평가 여부를
# 가중 평균해 ETF 자체의 '가치 점수'를 산출합니다.
#
# [동작 순서]
# 1. 4_financial_scanner/output 최신 xlsx → 종목코드별 상승여력 맵 구성
# 2. FDR로 국내 상장 ETF 전체 목록 수집
# 3. pykrx로 ETF별 구성 종목(PDF) 조회 → 비중 가중 평균 상승여력(가치점수) 계산
# 4. Gemini AI 배치 호출로 ETF 테마 자동 분류
# 5. 종합대시보드 + 테마별 시트 멀티시트 엑셀 저장
# 6. 분석 완료 후 이메일 발송 (SEND_EMAIL=true 시)
#
# [필요 패키지]
# pip install FinanceDataReader pykrx google-genai python-dotenv openpyxl pandas rich
#
# [환경 변수 (.env)]
# GEMINI_API_KEY=...
# KRX_ID=...  KRX_PW=...   (pykrx PDF 조회용)
# GMAIL_USER=...  GMAIL_APP_PASSWORD=...  NOTIFY_EMAIL=...
# =============================================================================

import os
import re
import sys
import glob
import json
import time
import requests
from concurrent.futures import ThreadPoolExecutor, as_completed

if sys.stdout.encoding and sys.stdout.encoding.lower() != "utf-8":
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")

from datetime import datetime, timezone, timedelta
import pandas as pd
from dotenv import load_dotenv

# .env는 pykrx 사용 전에 반드시 로드 (KRX_ID/KRX_PW 환경변수 필요)
load_dotenv()

from pykrx import stock as krx
import FinanceDataReader as fdr
from google import genai
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from rich.console import Console
from rich.progress import (
    BarColumn, MofNCompleteColumn, Progress, SpinnerColumn,
    TextColumn, TimeElapsedColumn, TimeRemainingColumn,
)

console = Console(legacy_windows=False)

GEMINI_API_KEY = os.getenv("GEMINI_API_KEY")
gemini_client   = genai.Client(api_key=GEMINI_API_KEY)

GMAIL_USER    = os.getenv("GMAIL_USER", "")
GMAIL_APP_PW  = os.getenv("GMAIL_APP_PASSWORD", "")
SEND_EMAIL    = os.getenv("SEND_EMAIL", "true").lower() in ("true", "1", "yes")
_extra_emails = [e.strip() for e in re.split(r"[,;]", os.getenv("NOTIFY_EMAILS_EXTRA", "")) if e.strip()]
_main_emails  = [e for e in [os.getenv("NOTIFY_EMAIL", ""), os.getenv("NOTIFY_EMAIL_WIFE", "")] if e.strip()]
NOTIFY_EMAILS = _main_emails + _extra_emails

KST = timezone(timedelta(hours=9))

def _desc(label: str, name: str, cols: int = 32) -> str:
    """터미널 표시 너비 기준으로 고정 폭 description 생성 (한글 2칸, ASCII 1칸)."""
    text = f"{label} {name}"
    w = sum(2 if ord(c) > 127 else 1 for c in text)
    return text + " " * max(0, cols - w)


THEME_LIST = [
    "반도체", "AI/IT", "배당", "채권", "원자재/금",
    "금융", "바이오/헬스케어", "소비재", "에너지",
    "부동산/리츠", "글로벌/해외", "기타",
]


# ─────────────────────────────────────────────────────────────────────────────
# S-RIM 데이터 로드
# ─────────────────────────────────────────────────────────────────────────────
def _load_srim_map() -> dict:
    """
    4_financial_scanner/output 최신 xlsx 종합 시트 → 티커(6자리) 기준 상승여력 맵 반환
    반환: {티커: {"종목명": ..., "상승여력1": float, "상승여력2": float, "종합점수": float}}
    """
    base  = os.path.join(os.path.dirname(__file__), "..", "4_financial_scanner", "output")
    files = sorted([f for f in glob.glob(os.path.join(base, "*.xlsx")) if "_hl" not in f])
    if not files:
        console.print("[yellow][WARN] 4_financial_scanner/output 파일 없음 — 가치점수 계산 생략[/yellow]")
        return {}

    df = pd.read_excel(files[-1], sheet_name="종합", dtype=str)

    def _f(row, col):
        try:
            return float(str(row.get(col, "")).replace(",", "").replace("%", "").strip())
        except Exception:
            return None

    result = {}
    for _, row in df.iterrows():
        ticker = str(row.get("티커", "")).strip().zfill(6)
        if not ticker or ticker == "000000":
            continue
        result[ticker] = {
            "종목명":      str(row.get("종목명", "")).strip(),
            "상승여력1":   _f(row, "상승여력1(%)"),
            "상승여력2":   _f(row, "상승여력2(%)"),
            "종합점수":    _f(row, "종합점수"),
            "PER":         _f(row, "PER"),
            "영업익증가율": _f(row, "영업익증가율(%)"),
        }

    console.print(f"  S-RIM 로드: {len(result)}개 종목 ({os.path.basename(files[-1])})")
    return result


# ─────────────────────────────────────────────────────────────────────────────
# ETF 목록 수집
# ─────────────────────────────────────────────────────────────────────────────
def _get_etf_list() -> pd.DataFrame:
    """FDR StockListing('ETF/KR')으로 국내 전체 ETF 목록 수집 및 컬럼 정규화"""
    df = fdr.StockListing("ETF/KR")
    rename = {}
    for c in df.columns:
        cl = c.lower().replace(" ", "")
        if cl in ("code", "symbol"):                   rename[c] = "티커"
        elif cl in ("name", "종목명"):                  rename[c] = "종목명"
        elif cl in ("marcap", "marketcap", "시가총액"):  rename[c] = "시가총액"
        elif "expense" in cl or "보수" in cl:           rename[c] = "보수율"
    df = df.rename(columns=rename)
    if "티커" in df.columns:
        df["티커"] = df["티커"].astype(str).str.zfill(6)
    return df.reset_index(drop=True)


# ─────────────────────────────────────────────────────────────────────────────
# ETF PDF(구성종목) 조회 → 가치점수 계산
# ─────────────────────────────────────────────────────────────────────────────
def _calc_value_score(etf_ticker: str, trade_date: str, srim_map: dict) -> dict:
    """
    pykrx로 ETF 구성 종목 조회 후 비중 가중 평균 상승여력 계산.
    반환: {"가치점수": float|None, "상승여력2": float|None, "매칭종목수": int, "구성종목수": int, "상위종목": str}
    """
    empty = {"가치점수": None, "상승여력2": None, "매칭종목수": 0, "구성종목수": 0, "상위종목": "", "분배율(%)": None}

    # 분배율(ttm) 조회
    div_yield = None
    try:
        url = f"https://m.stock.naver.com/api/stock/{etf_ticker}/integration"
        res = requests.get(url, headers={"User-Agent": "Mozilla/5.0"}, timeout=3)
        if res.status_code == 200:
            data = res.json()
            div_yield = data.get('etfKeyIndicator', {}).get('dividendYieldTtm')
    except Exception:
        pass

    try:
        pdf = krx.get_etf_portfolio_deposit_file(etf_ticker, trade_date)
        if pdf is None or pdf.empty:
            return empty

        pdf = pdf.reset_index()

        weight_col = next((c for c in pdf.columns if "비중" in str(c)), None)
        name_col   = next((c for c in pdf.columns if "종목명" in str(c) or str(c).lower() == "name"), None)
        ticker_col = next((c for c in pdf.columns
                           if str(c) in ("티커", "종목코드", "Code", "Symbol", "index", "Index")), None)

        if weight_col is None:
            return {**empty, "구성종목수": len(pdf)}

        # 비중 기준 내림차순 정렬을 위해 리스트에 수집
        parsed_rows = []
        for _, row in pdf.iterrows():
            try:
                w = float(str(row.get(weight_col, 0)).replace("%", "").replace(",", "") or 0)
            except Exception:
                w = 0.0
            if w <= 0:
                continue
            
            t = str(row.get(ticker_col, "")).strip().zfill(6) if ticker_col else ""
            n = str(row.get(name_col, "")) if name_col else ""
            parsed_rows.append((w, t, n))
            
        # 비중(w)이 큰 순서대로 정렬 후 상위 5개 추출
        parsed_rows.sort(key=lambda x: x[0], reverse=True)
        top5_rows = parsed_rows[:5]

        total_w1, weighted_up1 = 0.0, 0.0
        total_w2, weighted_up2 = 0.0, 0.0
        matched = 0

        for w, t, n in top5_rows:
            srim = srim_map.get(t)
            if srim:
                has_srim1 = srim.get("상승여력1") is not None
                has_srim2 = srim.get("상승여력2") is not None
                if has_srim1 or has_srim2:
                    matched += 1
                
                if has_srim1:
                    weighted_up1 += w * srim["상승여력1"]
                    total_w1     += w
                
                if has_srim2:
                    weighted_up2 += w * srim["상승여력2"]
                    total_w2     += w

        # 상위 5개 종목 내에서의 가중 평균 산출 (S-RIM 및 컨센서스)
        score1 = round(weighted_up1 / total_w1, 1) if total_w1 > 0 else None
        score2 = round(weighted_up2 / total_w2, 1) if total_w2 > 0 else None
        
        # 엑셀에 표시할 상위 종목명 (상위 5개)
        top_names = ", ".join(n for w, t, n in top5_rows if n)

        return {"가치점수": score1, "상승여력2": score2, "매칭종목수": matched, "구성종목수": len(pdf), "상위종목": top_names, "분배율(%)": div_yield}

    except Exception:
        return empty


# ─────────────────────────────────────────────────────────────────────────────
# Gemini 배치 테마 분류
# ─────────────────────────────────────────────────────────────────────────────
def _classify_themes(etf_rows: list[dict], batch_size: int = 40) -> dict:
    """
    ETF 이름 배치 전송 → Gemini JSON {티커: 테마} 반환
    etf_rows: [{"티커": ..., "종목명": ...}, ...]
    """
    result: dict[str, str] = {}
    total_batches = (len(etf_rows) + batch_size - 1) // batch_size

    for i in range(0, len(etf_rows), batch_size):
        batch     = etf_rows[i:i + batch_size]
        batch_num = i // batch_size + 1
        items_str = "\n".join(f'{r["티커"]} {r["종목명"]}' for r in batch)

        prompt = f"""아래 ETF 목록을 분석해 각 ETF에 가장 적합한 테마 하나를 JSON으로만 반환하세요.
테마 목록: {", ".join(THEME_LIST)}

ETF 목록:
{items_str}

응답 형식 (다른 텍스트 없이 JSON만):
{{"티커1": "테마명", "티커2": "테마명"}}"""

        for attempt in range(3):
            try:
                resp = gemini_client.models.generate_content(
                    model="gemini-2.5-flash-lite", contents=prompt
                )
                text = resp.text.strip()
                if text.startswith("```"):
                    text = text.split("```")[1]
                    if text.startswith("json"):
                        text = text[4:]
                result.update(json.loads(text.strip()))
                console.print(f"  테마분류 배치 {batch_num}/{total_batches} 완료")
                break
            except Exception as e:
                err = str(e)
                if "429" in err and attempt < 2:
                    time.sleep(30)
                elif ("503" in err or "UNAVAILABLE" in err) and attempt < 2:
                    time.sleep(10 * (2 ** attempt))
                else:
                    console.print(f"[yellow][WARN] 테마분류 배치 {batch_num} 실패: {e}[/yellow]")
                    break
        time.sleep(1)

    return result


# ─────────────────────────────────────────────────────────────────────────────
# 엑셀 서식 적용
# ─────────────────────────────────────────────────────────────────────────────
def _format_excel(path: str):
    HEADER_FILL  = PatternFill("solid", fgColor="1F4E79")
    HEADER_FONT  = Font(bold=True, color="FFFFFF")
    GREEN_FILL   = PatternFill("solid", fgColor="E2EFDA")   # 가치점수 >= 20
    YELLOW_FILL  = PatternFill("solid", fgColor="FFEB9C")   # 가치점수 >= 10

    def _cw(val):
        return sum(2 if ord(c) > 127 else 1 for c in str(val or ""))

    wb = openpyxl.load_workbook(path)
    for ws in wb.worksheets:
        for cell in ws[1]:
            cell.fill      = HEADER_FILL
            cell.font      = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", wrap_text=True)

        headers    = [c.value for c in ws[1]]
        score_col  = next((i+1 for i, h in enumerate(headers) if h == "가치점수"),  None)
        score2_col = next((i+1 for i, h in enumerate(headers) if h == "상승여력2"), None)
        ticker_col = next((i+1 for i, h in enumerate(headers) if h == "티커"),      None)
        name_col   = next((i+1 for i, h in enumerate(headers) if h == "종목명"),    None)

        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

            if score_col:
                sc = row[score_col - 1]
                try:
                    v = float(sc.value)
                    sc.fill = GREEN_FILL if v >= 20 else (YELLOW_FILL if v >= 10 else sc.fill)
                except Exception:
                    pass

            if score2_col:
                sc2 = row[score2_col - 1]
                try:
                    v2 = float(sc2.value)
                    sc2.fill = GREEN_FILL if v2 >= 20 else (YELLOW_FILL if v2 >= 10 else sc2.fill)
                except Exception:
                    pass

            if ticker_col and name_col:
                tc = row[ticker_col - 1]
                nc = row[name_col - 1]
                if tc.value and str(tc.value).strip():
                    code6 = str(tc.value).strip().zfill(6)
                    nc.hyperlink = f"https://finance.naver.com/item/main.naver?code={code6}"
                    nc.font = Font(color="0563C1", underline="single")

        for col in ws.columns:
            w = max((_cw(cell.value) for cell in col), default=8) + 2
            ws.column_dimensions[col[0].column_letter].width = min(w, 52)
        ws.freeze_panes = "A2"

    wb.save(path)


# ─────────────────────────────────────────────────────────────────────────────
# 이메일 발송
# ─────────────────────────────────────────────────────────────────────────────
def _send_email(output_file: str, df: pd.DataFrame):
    import smtplib
    from email.mime.application import MIMEApplication
    from email.mime.multipart import MIMEMultipart
    from email.mime.text import MIMEText
    from pathlib import Path

    if not (GMAIL_USER and GMAIL_APP_PW):
        console.print("[yellow]발송 계정 미설정[/yellow]")
        return
    if not NOTIFY_EMAILS:
        console.print("[yellow]수신자 미설정[/yellow]")
        return

    ts       = datetime.now(KST).strftime("%Y-%m-%d %H:%M")
    n_total  = len(df)
    n_scored = int(df["가치점수"].notna().sum()) if "가치점수" in df.columns else 0

    # 가치점수 상위 5개
    top5_lines = []
    if "가치점수" in df.columns and n_scored > 0:
        for _, r in df.nlargest(5, "가치점수").iterrows():
            top5_lines.append(f"  {str(r.get('종목명','')):<28} [{r.get('테마','')}]  {r.get('가치점수','')}")
    else:
        top5_lines = ["  (데이터 없음)"]

    # 테마별 1위 ETF
    theme_lines = []
    if "테마" in df.columns and "가치점수" in df.columns:
        for theme in THEME_LIST:
            grp = df[df["테마"] == theme].dropna(subset=["가치점수"])
            if grp.empty:
                continue
            best = grp.nlargest(1, "가치점수").iloc[0]
            theme_lines.append(f"  [{theme}] {best.get('종목명','')} (가치점수: {best.get('가치점수','')})")
    if not theme_lines:
        theme_lines = ["  (데이터 없음)"]

    body = "\n".join([
        "안녕하세요,",
        "",
        f"[{ts}] ETF 가치 분석 스캐너 분석이 완료되었습니다.",
        "",
        "=" * 60,
        "■ 분석 결과 요약",
        "=" * 60,
        f"  분석 ETF 수  : {n_total}개",
        f"  가치점수 산출: {n_scored}개 (구성종목 PDF 매칭 성공)",
        "",
        "■ 가치점수 상위 5개 ETF",
        *top5_lines,
        "",
        "■ 테마별 가치점수 1위 ETF",
        *theme_lines,
        "",
        "=" * 60,
        "■ 첨부 엑셀 구성",
        "=" * 60,
        "  [종합대시보드]",
        "    전체 ETF 가치점수 내림차순 정렬",
        "    컬럼: 티커 / 종목명 / 테마 / 시가총액 / 보수율 / 분배율(%)",
        "          가치점수 / 상승여력2 / 매칭종목수 / 구성종목수 / 상위종목",
        "    · 가치점수 및 상승여력2 >= 20% → 연초록",
        "    · 가치점수 >= 10% → 연노랑",
        "    · 종목명 클릭 → 네이버 금융 이동",
        "",
        "  [테마별 시트]",
        f"    {' / '.join(THEME_LIST)}",
        "    각 테마 내 가치점수 내림차순 정렬",
        "",
        "=" * 60,
        "■ 동작 방식",
        "=" * 60,
        "  FDR → 국내 전체 ETF 목록 수집",
        "  pykrx → ETF별 구성종목(PDF) 조회",
        "  4_financial_scanner S-RIM 데이터 → 비중 가중 평균 상승여력(가치점수) 산출",
        "  Gemini AI (gemini-2.5-flash-lite) → ETF 테마 자동 분류",
        "",
        "상세 분석 결과는 첨부 Excel 파일을 확인해주세요.",
        "",
        "※ 이 메일은 자동 발송입니다.",
    ])

    msg = MIMEMultipart()
    msg["From"]    = GMAIL_USER
    msg["To"]      = ", ".join(_main_emails) if _main_emails else GMAIL_USER
    if _extra_emails:
        msg["Bcc"] = ", ".join(_extra_emails)
    msg["Subject"] = f"[ETF 가치분석] {ts} — 가치점수 산출 {n_scored}개"
    msg.attach(MIMEText(body, "plain", "utf-8"))

    p = Path(output_file)
    if p.exists():
        with open(p, "rb") as f:
            part = MIMEApplication(f.read(), Name=p.name)
            part["Content-Disposition"] = f'attachment; filename="{p.name}"'
            msg.attach(part)

    try:
        with smtplib.SMTP_SSL("smtp.gmail.com", 465) as smtp:
            smtp.login(GMAIL_USER, GMAIL_APP_PW)
            smtp.sendmail(GMAIL_USER, NOTIFY_EMAILS, msg.as_string())
        _to_str  = ', '.join(_main_emails) if _main_emails else GMAIL_USER
        _bcc_str = f" | Bcc: {len(_extra_emails)}명" if _extra_emails else ""
        console.print(f"  [green]이메일 발송 완료 — To: {_to_str}{_bcc_str}[/green]")
        if _extra_emails:
            console.print(f"  [dim]  Bcc: {', '.join(_extra_emails)}[/dim]")
    except Exception as e:
        console.print(f"  [red]이메일 발송 실패: {e}[/red]")


# ─────────────────────────────────────────────────────────────────────────────
# 메인 실행
# ─────────────────────────────────────────────────────────────────────────────
console.print("[bold cyan]═══ ETF 가치 분석 스캐너 ═══[/bold cyan]")

# 1. S-RIM 맵 로드
srim_map = _load_srim_map()

# 2. ETF 목록 수집
console.print("\nETF 목록 수집 중...")
df_etf = _get_etf_list()
console.print(f"  총 {len(df_etf)}개 ETF 수집 완료")

# 3. 최근 거래일 (pykrx PDF 조회 기준)
try:
    trade_date = krx.get_nearest_business_day_in_a_week()
except Exception:
    trade_date = (datetime.now(KST) - timedelta(days=1)).strftime("%Y%m%d")
console.print(f"  PDF 기준일: {trade_date}")

# 4. PDF 수집 + 가치점수 계산 (병렬 처리)
etf_list   = [(str(row.get("티커", "")).strip().zfill(6), str(row.get("종목명", "")))
              for _, row in df_etf.iterrows()]
value_rows = [None] * len(etf_list)

with Progress(
    SpinnerColumn(),
    TextColumn("{task.description}"),
    BarColumn(),
    MofNCompleteColumn(),
    TimeElapsedColumn(),
    TimeRemainingColumn(),
    console=console,
    transient=False,
) as progress:
    task = progress.add_task("PDF 수집 및 가치점수 계산...", total=len(etf_list))
    with ThreadPoolExecutor(max_workers=8) as executor:
        futures = {
            executor.submit(_calc_value_score, ticker, trade_date, srim_map): idx
            for idx, (ticker, _) in enumerate(etf_list)
        }
        for future in as_completed(futures):
            idx  = futures[future]
            name = etf_list[idx][1]
            value_rows[idx] = future.result()
            progress.update(task, description=_desc("PDF:", name[:18]))
            progress.advance(task)

# 5. 결과 병합
df_scores = pd.DataFrame(value_rows)
df_result = pd.concat([df_etf.reset_index(drop=True), df_scores], axis=1)

# 6. Gemini 테마 분류
console.print("\nGemini 테마 분류 중...")
etf_rows  = [{"티커": str(r.get("티커", "")), "종목명": str(r.get("종목명", ""))}
             for _, r in df_result.iterrows()]
theme_map = _classify_themes(etf_rows)
df_result["테마"] = df_result["티커"].map(theme_map).fillna("기타")

# 컬럼 순서 정리 + 가치점수 내림차순 정렬
FRONT = ["티커", "종목명", "테마", "시가총액", "보수율", "분배율(%)", "가치점수", "상승여력2", "매칭종목수", "구성종목수", "상위종목"]
others = [c for c in df_result.columns if c not in FRONT]
df_result = df_result[[c for c in FRONT if c in df_result.columns] + others]
df_result = df_result.sort_values("가치점수", ascending=False, na_position="last").reset_index(drop=True)

# 7. 엑셀 저장
output_dir  = os.path.join(os.path.dirname(os.path.abspath(__file__)), "output")
os.makedirs(output_dir, exist_ok=True)
ts          = datetime.now(KST).strftime("%Y_%m%d_%H%M")
output_file = os.path.join(output_dir, f"ETF_가치분석_{ts}.xlsx")

with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
    df_result.to_excel(writer, sheet_name="종합대시보드", index=False)
    for theme in THEME_LIST:
        df_t = df_result[df_result["테마"] == theme]
        if not df_t.empty:
            sheet_name = theme.replace("/", "·")[:31]
            df_t.to_excel(writer, sheet_name=sheet_name, index=False)

try:
    import importlib.util
    _spec = importlib.util.spec_from_file_location(
        "format_excel",
        os.path.join(os.path.dirname(os.path.abspath(__file__)), "format_excel.py")
    )
    _mod = importlib.util.module_from_spec(_spec)
    _spec.loader.exec_module(_mod)
    _mod.format_file(output_file)
    fmt_path = os.path.splitext(output_file)[0] + "_fmt.xlsx"
    import shutil
    shutil.move(fmt_path, output_file)
except Exception as e:
    console.print(f"[yellow][WARN] 상세 포맷 실패: {e}[/yellow]")
    _format_excel(output_file)

console.print(f"\n[bold green][OK] 분석 완료![/bold green] {output_file}")

# 8. 이메일 발송
if SEND_EMAIL:
    _send_email(output_file, df_result)
else:
    console.print("[dim]이메일 발송 생략 (SEND_EMAIL=false)[/dim]")
