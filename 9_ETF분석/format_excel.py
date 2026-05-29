"""
ETF 가치 분석 엑셀 포맷터
- 종합대시보드 + 테마별 전체 시트에 헤더 스타일 적용
- 가치점수 조건부 색상 (>=20 연초록, >=10 연노랑)
- 종목명 → 네이버 금융 하이퍼링크
- xlwings AutoFit (행 높이 자동 조정)
사용법: python format_excel.py [파일경로]  (생략 시 output 최신 파일 자동 선택)
"""
import sys
import glob
import os
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

# ── 공통 스타일 ──────────────────────────────────────────────────────────────
HEADER_FILL  = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT  = Font(color="FFFFFF", bold=True, size=10)
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
THIN         = Side(style="thin", color="BBBBBB")
BORDER       = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

GREEN_FILL  = PatternFill("solid", fgColor="C6EFCE")   # 가치점수 >= 20
YELLOW_FILL = PatternFill("solid", fgColor="FFEB9C")   # 가치점수 >= 10

# ── 컬럼 너비 ────────────────────────────────────────────────────────────────
COLUMN_WIDTHS = {
    "티커":      8,
    "종목명":    32,
    "테마":      12,
    "시가총액":  16,
    "보수율":    8,
    "가치점수":  10,
    "매칭종목수": 10,
    "구성종목수": 10,
    "상위종목":  40,
}

WRAP_COLS = {"상위종목"}


def _style_sheet(ws) -> None:
    """단일 시트에 전체 서식 적용."""
    headers = [c.value for c in ws[1]]
    col_idx = {h: i + 1 for i, h in enumerate(headers) if h}

    # 헤더 행
    ws.row_dimensions[1].height = 22
    for cell in ws[1]:
        cell.fill      = HEADER_FILL
        cell.font      = HEADER_FONT
        cell.alignment = HEADER_ALIGN
        cell.border    = BORDER

    # 컬럼 너비
    for h, idx in col_idx.items():
        ws.column_dimensions[get_column_letter(idx)].width = COLUMN_WIDTHS.get(h, 14)

    score_col  = col_idx.get("가치점수")
    name_col   = col_idx.get("종목명")
    ticker_col = col_idx.get("티커")

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        ws.row_dimensions[row[0].row].height = 18

        for cell in row:
            h    = headers[cell.column - 1] if cell.column <= len(headers) else ""
            wrap = h in WRAP_COLS
            cell.border    = BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=wrap, shrink_to_fit=not wrap)

        # 가치점수 조건부 색상
        if score_col:
            sc = row[score_col - 1]
            try:
                v = float(sc.value)
                if v >= 20:
                    sc.fill = GREEN_FILL
                elif v >= 10:
                    sc.fill = YELLOW_FILL
            except Exception:
                pass

        # 종목명 → 네이버 금융 하이퍼링크
        if ticker_col and name_col:
            tc = row[ticker_col - 1]
            nc = row[name_col - 1]
            if tc.value and str(tc.value).strip():
                code6 = str(tc.value).strip().zfill(6)
                nc.hyperlink = f"https://finance.naver.com/item/main.naver?code={code6}"
                nc.font = Font(color="0563C1", underline="single", bold=True)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def format_file(path: str):
    wb = openpyxl.load_workbook(path)
    for ws in wb.worksheets:
        _style_sheet(ws)

    base, ext = os.path.splitext(path)
    out_path  = f"{base}_fmt{ext}"
    wb.save(out_path)

    # xlwings AutoFit (행 높이 자동 조정)
    try:
        import xlwings as xw
        with xw.App(visible=False) as app:
            wb2 = app.books.open(os.path.abspath(out_path))
            for sheet in wb2.sheets:
                sheet.autofit("rows")
                last_row = sheet.used_range.last_cell.row
                for r in range(2, last_row + 1):
                    h = sheet.range(f"A{r}").row_height
                    if h and h > 15:
                        sheet.range(f"A{r}").row_height = h + 4
            wb2.save()
            wb2.close()
        print("[OK] 행 높이 자동 조정 완료")
    except Exception as e:
        print(f"[WARN] AutoFit 실패 (xlwings): {e}")

    print(f"[OK] 포맷 완료: {out_path}")


if __name__ == "__main__":
    if len(sys.argv) > 1:
        target = sys.argv[1]
    else:
        base  = os.path.join(os.path.dirname(os.path.abspath(__file__)), "output")
        files = sorted([f for f in glob.glob(os.path.join(base, "ETF_가치분석_*.xlsx"))
                        if "~$" not in f and "_fmt" not in f])
        if not files:
            print("[ERR] output 폴더에 xlsx 파일이 없습니다.")
            sys.exit(1)
        target = files[-1]
        print(f"대상 파일: {target}")
    format_file(target)
