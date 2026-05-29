"""
================================================
  Markdown + Excel 리포트 생성기
  매일 스크리닝 결과를 report/output/YYYY-MM-DD.md / .xlsx 로 저장
  완료 후 회사 이메일로 자동 발송 (Gmail SMTP_SSL)
================================================
"""

import smtplib
from datetime import date
from email import encoders
from email.mime.base import MIMEBase
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from pathlib import Path
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from config import (
    ACTUAL_YEAR, ESTIMATE_YEAR,
    MIN_COVERAGE, MIN_MARKET_CAP,
    MIN_OP_PROFIT_GROWTH, MIN_OPM_IMPROVEMENT, MIN_REVENUE_GROWTH,
    REPORT_DIR,
    EMAIL_USER, EMAIL_PASS, NOTIFY_EMAILS, MAIN_EMAILS, EXTRA_EMAILS,
)
from crawler.fnguide import ConsensusData
from screener.scorer import Score


# ── 스타일 상수 ───────────────────────────────────────────────────────────────

_HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
_HEADER_FONT = Font(bold=True, color="FFFFFF", size=10)
_LINK_FONT   = Font(color="0563C1", underline="single", bold=True, size=10)
_BASE_FONT   = Font(size=10)
_CENTER      = Alignment(horizontal="center", vertical="center", wrap_text=False)
_LEFT        = Alignment(horizontal="left",   vertical="center")


def _col_width(text: str) -> float:
    """한글·영문 혼합 텍스트의 대략적 열 너비 계산."""
    w = 0.0
    for ch in str(text):
        w += 2.0 if ord(ch) > 127 else 1.0
    return max(w + 2, 6)


def _auto_width(ws):
    """각 열의 최대 콘텐츠 기준으로 열 너비 자동 조정."""
    for col_cells in ws.columns:
        max_w = 6.0
        for cell in col_cells:
            if cell.value is not None:
                max_w = max(max_w, _col_width(cell.value))
        ws.column_dimensions[get_column_letter(col_cells[0].column)].width = min(max_w, 40)


def _fmt(value: Optional[float], suffix: str = "", decimals: int = 1) -> str:
    if value is None:
        return "N/A"
    return f"{value:,.{decimals}f}{suffix}"


def _arrow(value: Optional[float]) -> str:
    if value is None:
        return ""
    return "▲" if value > 0 else ("▼" if value < 0 else "─")


# ── 이메일 ────────────────────────────────────────────────────────────────────

def send_email(excel_path: Path, md_path: Path, run_date: date, count: int):
    """Gmail SMTP_SSL로 Excel + Markdown 리포트를 회사 메일로 발송."""
    if not all([EMAIL_USER, EMAIL_PASS, NOTIFY_EMAILS]):
        print("  이메일 설정 미완료 (.env 확인): GMAIL_USER / GMAIL_APP_PASSWORD / NOTIFY_EMAIL")
        return

    msg = MIMEMultipart()
    msg["From"]    = EMAIL_USER
    msg["To"]      = ", ".join(MAIN_EMAILS) if MAIN_EMAILS else EMAIL_USER
    if EXTRA_EMAILS:
        msg["Bcc"] = ", ".join(EXTRA_EMAILS)
    msg["Subject"] = f"[이익 스크리너] {run_date} 스크리닝 결과 — {count}개 통과"

    body = (
        f"안녕하세요,\n\n"
        f"{run_date} 기준 이익 추정치 급등 종목 스크리닝 결과를 전달드립니다.\n\n"
        f"  · 통과 종목: {count}개\n"
        f"  · 상세 내용은 첨부 Excel 파일을 확인해 주세요.\n\n"
        f"[5축 필터 조건]\n"
        f"  1) 이익 모멘텀: 매출 ≥{MIN_REVENUE_GROWTH:.0f}%, 영익 ≥{MIN_OP_PROFIT_GROWTH:.0f}%(or 흑전), OPM ≥{MIN_OPM_IMPROVEMENT:.0f}%p\n"
        f"  2) 컨센 모멘텀: 1M·3M 컨센 상향\n"
        f"  3) 밸류에이션: Fwd PER < 30, 52주고점대비 < 90%\n"
        f"  4) 리스크 제거: 관리종목 제외, 부채비율 < 300%\n"
        f"  5) 수급: 외국인/기관 3M 순매수 > 0\n\n"
        f"자동 생성됨 | 데이터 출처: FnGuide / 네이버금융 / DART / KRX"
    )
    msg.attach(MIMEText(body, "plain", "utf-8"))

    for path in [excel_path, md_path]:
        if path.exists():
            with open(path, "rb") as f:
                part = MIMEBase("application", "octet-stream")
                part.set_payload(f.read())
            encoders.encode_base64(part)
            part.add_header("Content-Disposition", f'attachment; filename="{path.name}"')
            msg.attach(part)

    try:
        with smtplib.SMTP_SSL("smtp.gmail.com", 465) as smtp:
            smtp.login(EMAIL_USER, EMAIL_PASS)
            smtp.sendmail(EMAIL_USER, NOTIFY_EMAILS, msg.as_string())
        _to_str  = ', '.join(MAIN_EMAILS) if MAIN_EMAILS else EMAIL_USER
        _bcc_str = f" | Bcc: {len(EXTRA_EMAILS)}명" if EXTRA_EMAILS else ""
        print(f"  이메일 발송 완료 — To: {_to_str}{_bcc_str}")
        if EXTRA_EMAILS:
            print(f"    Bcc: {', '.join(EXTRA_EMAILS)}")
    except Exception as e:
        print(f"  이메일 발송 실패: {e}")


# ── Excel 생성 ────────────────────────────────────────────────────────────────

def _build_excel(ranked: list[tuple[ConsensusData, Score]], run_date: date) -> Path:
    """스크리닝 결과를 스타일 적용된 Excel로 저장."""
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    excel_path = REPORT_DIR / f"{run_date}.xlsx"

    wb = Workbook()

    # ── 시트 1: 요약 ──────────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "스크리닝 요약"
    ws.freeze_panes = "A2"

    headers = [
        "순위", "종목명", "코드",
        f"매출({ACTUAL_YEAR}A→{ESTIMATE_YEAR}E)",
        f"영업이익({ACTUAL_YEAR}A→{ESTIMATE_YEAR}E)",
        "OPM 개선", "컨센 1M 상향", "종합점수",
    ]
    ws.append(headers)

    # 헤더 스타일
    for col_idx, _ in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill      = _HEADER_FILL
        cell.font      = _HEADER_FONT
        cell.alignment = _CENTER

    # 데이터 행
    for rank, (data, sc) in enumerate(ranked, 1):
        op_str = "흑자전환" if data.turnaround else _fmt(data.op_profit_growth, "%")
        row_data = [
            rank,
            data.name,
            data.code,
            _fmt(data.revenue_growth, "%"),
            op_str,
            _fmt(data.opm_improvement, "%p"),
            _fmt(data.op_revision_1m, "%"),
            round(sc.total, 1),
        ]
        ws.append(row_data)
        row_idx = rank + 1

        # 종목명 하이퍼링크
        name_cell = ws.cell(row=row_idx, column=2)
        name_cell.hyperlink = f"https://finance.naver.com/item/main.naver?code={data.code}"
        name_cell.font      = _LINK_FONT

        # 숫자 포맷 및 정렬
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.alignment = _CENTER
            if col_idx == 2:
                cell.alignment = _LEFT
            if col_idx == 8:
                cell.number_format = "0.0"

    _auto_width(ws)
    ws.row_dimensions[1].height = 22

    # ── 시트 2: 상세 ──────────────────────────────────────────────────────────
    ws2 = wb.create_sheet("상세 지표")
    ws2.freeze_panes = "A2"

    detail_headers = [
        "순위", "종목명", "코드", "시총(억)", "커버리지",
        f"매출 {ACTUAL_YEAR}A(억)", f"매출 {ESTIMATE_YEAR}E(억)", "매출증가",
        f"영익 {ACTUAL_YEAR}A(억)", f"영익 {ESTIMATE_YEAR}E(억)", "영익증가",
        f"OPM {ACTUAL_YEAR}A(%)", f"OPM {ESTIMATE_YEAR}E(%)", "OPM개선",
        "컨센 1M", "컨센 3M", "종합점수",
    ]
    ws2.append(detail_headers)

    for col_idx, _ in enumerate(detail_headers, 1):
        cell = ws2.cell(row=1, column=col_idx)
        cell.fill      = _HEADER_FILL
        cell.font      = _HEADER_FONT
        cell.alignment = _CENTER

    for rank, (data, sc) in enumerate(ranked, 1):
        op_str = "흑자전환" if data.turnaround else _fmt(data.op_profit_growth, "%")
        row_data = [
            rank,
            data.name,
            data.code,
            round(data.market_cap) if data.market_cap else None,
            data.analyst_count,
            round(data.revenue_25a) if data.revenue_25a else None,
            round(data.revenue_26e) if data.revenue_26e else None,
            _fmt(data.revenue_growth, "%"),
            round(data.op_profit_25a) if data.op_profit_25a else None,
            round(data.op_profit_26e) if data.op_profit_26e else None,
            op_str,
            _fmt(data.opm_25a, "%"),
            _fmt(data.opm_26e, "%"),
            _fmt(data.opm_improvement, "%p"),
            _fmt(data.op_revision_1m, "%"),
            _fmt(data.op_revision_3m, "%"),
            round(sc.total, 1),
        ]
        ws2.append(row_data)
        row_idx = rank + 1

        name_cell = ws2.cell(row=row_idx, column=2)
        name_cell.hyperlink = f"https://finance.naver.com/item/main.naver?code={data.code}"
        name_cell.font      = _LINK_FONT

        for col_idx in range(1, len(detail_headers) + 1):
            cell = ws2.cell(row=row_idx, column=col_idx)
            cell.alignment = _CENTER
            if col_idx == 2:
                cell.alignment = _LEFT

    _auto_width(ws2)
    ws2.row_dimensions[1].height = 22

    wb.save(excel_path)
    return excel_path


# ── Markdown 생성 ─────────────────────────────────────────────────────────────

def generate(ranked: list[tuple[ConsensusData, Score]], run_date: date = None) -> Path:
    """스크리닝 결과를 Markdown + Excel로 저장하고 이메일 발송."""
    today = run_date or date.today()
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    md_path = REPORT_DIR / f"{today}.md"

    lines = [
        f"# 이익 추정치 급등 종목 스크리닝 — {today}",
        "",
        "## 스크리닝 조건",
        "",
        f"| 항목 | 기준 |",
        f"|------|------|",
        f"| 유니버스 | KOSPI+KOSDAQ, 시총 {MIN_MARKET_CAP:,}억 이상, 커버리지 {MIN_COVERAGE}개 이상 |",
        f"| 매출 증가율 | {ACTUAL_YEAR}A → {ESTIMATE_YEAR}E ≥ {MIN_REVENUE_GROWTH:.0f}% |",
        f"| 영업이익 증가율 | {ACTUAL_YEAR}A → {ESTIMATE_YEAR}E ≥ {MIN_OP_PROFIT_GROWTH:.0f}% 또는 흑자전환 |",
        f"| OPM 개선폭 | {ESTIMATE_YEAR}E − {ACTUAL_YEAR}A ≥ {MIN_OPM_IMPROVEMENT:.0f}%p |",
        f"| 컨센서스 방향 | 최근 1개월 {ESTIMATE_YEAR}E 영업이익 컨센 상향 |",
        "",
        f"**통과 종목: {len(ranked)}개**",
        "",
    ]

    if not ranked:
        lines.append("> 조건을 충족하는 종목이 없습니다.")
        md_path.write_text("\n".join(lines), encoding="utf-8")
        _build_excel(ranked, today)
        return md_path

    lines += [
        "## 통과 종목 요약",
        "",
        "| 순위 | 종목명 | 종목코드 | 매출증가율 | 영업이익증가율 | OPM개선 | 컨센상향(1M) | 점수 |",
        "|------|--------|----------|-----------|--------------|---------|------------|------|",
    ]

    for rank, (data, sc) in enumerate(ranked, 1):
        op_growth_str = "흑자전환" if data.turnaround else _fmt(data.op_profit_growth, "%")
        lines.append(
            f"| {rank} "
            f"| {data.name} "
            f"| {data.code} "
            f"| {_arrow(data.revenue_growth)} {_fmt(data.revenue_growth, '%')} "
            f"| {_arrow(data.op_profit_growth)} {op_growth_str} "
            f"| {_arrow(data.opm_improvement)} {_fmt(data.opm_improvement, '%p')} "
            f"| {_arrow(data.op_revision_1m)} {_fmt(data.op_revision_1m, '%')} "
            f"| **{sc.total:.1f}** |"
        )

    lines += [
        "",
        "## 종목별 상세",
        "",
    ]

    for rank, (data, sc) in enumerate(ranked, 1):
        op_growth_str = "**흑자전환**" if data.turnaround else _fmt(data.op_profit_growth, "%")
        lines += [
            f"### {rank}. {data.name} ({data.code})",
            "",
            f"- 시가총액: {_fmt(data.market_cap, '억', 0)}  |  커버리지: {data.analyst_count}개 증권사",
            f"- 종합점수: **{sc.total:.1f}점**",
            "",
            f"| 지표 | {ACTUAL_YEAR}A | {ESTIMATE_YEAR}E | 변화 |",
            f"|------|-------|-------|------|",
            f"| 매출액 (억원) | {_fmt(data.revenue_25a, '억', 0)} | {_fmt(data.revenue_26e, '억', 0)} | {_arrow(data.revenue_growth)} {_fmt(data.revenue_growth, '%')} |",
            f"| 영업이익 (억원) | {_fmt(data.op_profit_25a, '억', 0)} | {_fmt(data.op_profit_26e, '억', 0)} | {_arrow(data.op_profit_growth)} {op_growth_str} |",
            f"| 영업이익률 (%) | {_fmt(data.opm_25a, '%')} | {_fmt(data.opm_26e, '%')} | {_arrow(data.opm_improvement)} {_fmt(data.opm_improvement, '%p')} |",
            f"| 컨센 상향 (1M) | — | — | {_arrow(data.op_revision_1m)} {_fmt(data.op_revision_1m, '%')} |",
            "",
        ]

    lines += [
        "---",
        f"*생성일시: {today}  |  데이터 출처: FnGuide / 네이버금융 / DART*",
    ]

    md_path.write_text("\n".join(lines), encoding="utf-8")

    # Excel 생성 + 이메일 발송
    excel_path = _build_excel(ranked, today)
    send_email(excel_path, md_path, today, len(ranked))

    return md_path
